import io
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from reportlab.pdfgen import canvas

from app.services.document_intake import sanitize_and_extract
from app.services.upload_policy import DocumentSafetyError, sanitize_filename, sanitize_upload


def _pdf(text="Invoice number: INV-123"):
    stream = io.BytesIO()
    pdf = canvas.Canvas(stream)
    pdf.drawString(40, 780, text)
    pdf.save()
    return stream.getvalue()


@pytest.mark.unit
def test_filename_and_signature_policy_blocks_traversal_and_mismatch():
    assert sanitize_filename(r"..\folder\ Invoice  1.pdf") == "Invoice 1.pdf"
    with pytest.raises(DocumentSafetyError) as caught:
        sanitize_upload("invoice.pdf", "text/plain", b"not a pdf")
    assert caught.value.code == "TYPE_MISMATCH"
    with pytest.raises(DocumentSafetyError) as caught:
        sanitize_upload("invoice.exe", "application/octet-stream", b"MZ")
    assert caught.value.code == "UNSUPPORTED_TYPE"


@pytest.mark.unit
def test_pdf_csv_and_image_extraction_are_bounded_and_no_bytes_retained():
    result = sanitize_and_extract("invoice.pdf", "application/pdf", _pdf())
    assert result.status == "extracted"
    assert result.normalized_fields["invoice_number"] == "INV-123"
    assert not hasattr(result, "data")

    csv = b"field,value\nformula,=SUM(A1:A2)\n"
    parsed = sanitize_and_extract("sheet.csv", "text/csv", csv)
    assert "'=SUM(A1:A2)" in parsed.text

    image = io.BytesIO()
    Image.new("RGB", (2, 2), "white").save(image, format="PNG")
    ocr = sanitize_and_extract("scan.png", "image/png", image.getvalue())
    assert ocr.status == "needs_ocr"
    assert ocr.error_code == "OCR_UNAVAILABLE"
    assert ocr.parser_provenance == "image_metadata_only"

    doc = Document()
    doc.add_paragraph("Purchase order: PO-456")
    docx_stream = io.BytesIO()
    doc.save(docx_stream)
    docx_result = sanitize_and_extract(
        "order.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        docx_stream.getvalue(),
    )
    assert docx_result.normalized_fields["purchase_order"] == "PO-456"

    workbook = Workbook()
    workbook.active.append(["Amount", "=SUM(A1:A2)"])
    xlsx_stream = io.BytesIO()
    workbook.save(xlsx_stream)
    xlsx_result = sanitize_and_extract(
        "sheet.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        xlsx_stream.getvalue(),
    )
    assert xlsx_result.status == "extracted"


@pytest.mark.unit
def test_malformed_office_zip_and_empty_or_oversized_uploads_fail_safely():
    with pytest.raises(DocumentSafetyError) as caught:
        sanitize_upload("data.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"PK\x03\x04bad")
    assert caught.value.code == "TYPE_MISMATCH"
    with pytest.raises(DocumentSafetyError) as caught:
        sanitize_upload("empty.pdf", "application/pdf", b"")
    assert caught.value.code == "EMPTY_FILE"


@pytest.mark.unit
def test_office_archive_rejects_path_traversal_metadata():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<types/>")
        archive.writestr("word/document.xml", "<doc/>")
        archive.writestr("../evil.txt", "x")
    with pytest.raises(DocumentSafetyError) as caught:
        sanitize_upload("doc.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", stream.getvalue())
    assert caught.value.code == "UNSAFE_ARCHIVE"
